from __future__ import annotations

from collections import Counter, defaultdict
from copy import copy
from datetime import datetime
from pathlib import Path
import re

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


RESPONSES = Path(r"C:\Users\riyav\Downloads\Mawla Ghati Ultra 2026 (Responses).xlsx")
GROUP = Path(r"D:\Teakwood\Mawla\Mawla Ghaati Group Booking.xlsx")
FINAL = Path(r"D:\Teakwood\Mawla\Mawla Ghaati Group Booking_Chennai1.xlsx")
OUT_DIR = Path(r"D:\Teakwood\Mawla\mawla-booking\outputs")
OUT_FILE = OUT_DIR / "Mawla Ghaati Group Booking_Chennai1_checked.xlsx"


KEY_COLS = {
    "name": "Full Name",
    "gender": "Gender",
    "phone": "Phone Number",
    "emergency": "Emergency Contact",
    "city": "City (travelling from)",
    "state": "State (travelling from)",
    "accommodation": "Accommodation (per night charges)",
    "date": "Date of Check In",
    "nights": "No of nights of stay",
}


def norm_text(value):
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value).strip()).casefold()


def norm_phone(value):
    if value is None:
        return ""
    text = str(value).strip()
    if text.endswith(".0"):
        text = text[:-2]
    digits = re.sub(r"\D", "", text)
    if len(digits) > 10 and digits.startswith("91"):
        digits = digits[-10:]
    return digits


def norm_date(value):
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    text = str(value).strip()
    for fmt in ("%m/%d/%Y", "%d/%m/%Y", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            pass
    return text


def rows_from_sheet(ws):
    headers = [cell.value for cell in ws[1]]
    index = {str(h).strip(): idx + 1 for idx, h in enumerate(headers) if h is not None}
    rows = []
    for row_num in range(2, ws.max_row + 1):
        if not any(ws.cell(row_num, col).value is not None for col in range(1, min(ws.max_column, 10) + 1)):
            continue
        rec = {"_row": row_num, "_sheet": ws.title}
        for key, col_name in KEY_COLS.items():
            col = index.get(col_name)
            if col is None and col_name == "Date of Check In":
                col = index.get("Date of Check In (MM/DD/YYYY)")
            if col is None and col_name == "No of nights of stay":
                col = index.get("No of nights of stay : 1")
            rec[key] = ws.cell(row_num, col).value if col else None
        rec["_name_key"] = norm_text(rec["name"])
        rec["_phone_key"] = norm_phone(rec["phone"])
        rec["_date_key"] = norm_date(rec["date"])
        rec["_acc_key"] = norm_text(rec["accommodation"])
        rows.append(rec)
    return rows


def copy_row_values_and_styles(ws, source_row, target_row):
    for col in range(1, ws.max_column + 1):
        src = ws.cell(source_row, col)
        dst = ws.cell(target_row, col)
        dst.value = src.value
        if src.has_style:
            dst._style = copy(src._style)
        if src.number_format:
            dst.number_format = src.number_format


def main():
    responses_wb = openpyxl.load_workbook(RESPONSES, data_only=True)
    source_chennai = rows_from_sheet(responses_wb["Sheet3"])
    source_chennai_by_name_phone = {(r["_name_key"], r["_phone_key"]): r for r in source_chennai}
    source_chennai_by_phone = defaultdict(list)
    for rec in source_chennai:
        source_chennai_by_phone[rec["_phone_key"]].append(rec)

    group_wb = openpyxl.load_workbook(GROUP, data_only=True)
    group_rows = rows_from_sheet(group_wb["Sheet1"])

    wb = openpyxl.load_workbook(FINAL)
    ws = wb["Sheet1"]
    final_rows = rows_from_sheet(ws)

    notes = []
    changed = []

    # The final sheet is separated into date bands by blank rows.
    date_by_row = {}
    current_date = datetime(2026, 6, 4)
    for row_num in range(2, ws.max_row + 1):
        has_data = any(ws.cell(row_num, col).value is not None for col in range(1, 10))
        if not has_data:
            if current_date.date().isoformat() == "2026-06-04":
                current_date = datetime(2026, 6, 5)
            elif current_date.date().isoformat() == "2026-06-05":
                current_date = datetime(2026, 6, 6)
            continue
        date_by_row[row_num] = current_date

    for rec in final_rows:
        row_num = rec["_row"]
        intended_date = date_by_row.get(row_num)
        if intended_date:
            old = ws.cell(row_num, 8).value
            if norm_date(old) != intended_date.date().isoformat():
                ws.cell(row_num, 8).value = intended_date
                ws.cell(row_num, 8).number_format = "yyyy-mm-dd"
                changed.append((row_num, "Date of Check In", old, intended_date.date().isoformat()))

        source = source_chennai_by_name_phone.get((rec["_name_key"], rec["_phone_key"]))
        if source is None and rec["_phone_key"] in source_chennai_by_phone and len(source_chennai_by_phone[rec["_phone_key"]]) == 1:
            candidate = source_chennai_by_phone[rec["_phone_key"]][0]
            if candidate["_name_key"] != rec["_name_key"]:
                notes.append((row_num, rec["name"], rec["phone"], "Name differs from source row", candidate["name"]))
            source = candidate

        if source:
            final_acc = norm_text(ws.cell(row_num, 7).value)
            if final_acc != source["_acc_key"]:
                old = ws.cell(row_num, 7).value
                ws.cell(row_num, 7).value = source["accommodation"]
                changed.append((row_num, "Accommodation", old, source["accommodation"]))
        else:
            notes.append((row_num, rec["name"], rec["phone"], "No matching Chennai source row", "Kept as-is"))

    final_rows_after = rows_from_sheet(ws)
    combo_counts = Counter(
        (r["_name_key"], r["_phone_key"], norm_date(ws.cell(r["_row"], 8).value), r["_acc_key"])
        for r in final_rows_after
    )
    phone_date_counts = Counter(
        (r["_phone_key"], norm_date(ws.cell(r["_row"], 8).value))
        for r in final_rows_after
        if r["_phone_key"]
    )
    name_date_counts = Counter(
        (r["_name_key"], norm_date(ws.cell(r["_row"], 8).value))
        for r in final_rows_after
        if r["_name_key"]
    )

    for rec in final_rows_after:
        date = norm_date(ws.cell(rec["_row"], 8).value)
        if combo_counts[(rec["_name_key"], rec["_phone_key"], date, rec["_acc_key"])] > 1:
            notes.append((rec["_row"], rec["name"], rec["phone"], "Exact duplicate row key", date))
        if phone_date_counts[(rec["_phone_key"], date)] > 1:
            notes.append((rec["_row"], rec["name"], rec["phone"], "Duplicate primary phone on same date", date))
        if name_date_counts[(rec["_name_key"], date)] > 1:
            notes.append((rec["_row"], rec["name"], rec["phone"], "Duplicate name on same date", date))

    final_name_phones = {(r["_name_key"], r["_phone_key"]) for r in final_rows_after}
    for rec in source_chennai:
        if (rec["_name_key"], rec["_phone_key"]) not in final_name_phones:
            notes.append(("", rec["name"], rec["phone"], "Chennai source row missing from final", rec["_date_key"]))

    for rec in group_rows:
        if (rec["_name_key"], rec["_phone_key"]) not in final_name_phones:
            notes.append(("", rec["name"], rec["phone"], "Group booking source row not present in Chennai final", rec["_date_key"]))

    if "Audit" in wb.sheetnames:
        del wb["Audit"]
    audit = wb.create_sheet("Audit")
    audit_rows = [
        ["Section", "Row", "Name", "Phone", "Field / Issue", "Old / Detail", "New / Source"],
    ]
    for row_num, field, old, new in changed:
        audit_rows.append(["Corrected", row_num, ws.cell(row_num, 1).value, ws.cell(row_num, 3).value, field, old, new])
    for row_num, name, phone, issue, detail in notes:
        audit_rows.append(["Review", row_num, name, phone, issue, detail, ""])

    for row in audit_rows:
        audit.append(row)
    audit.freeze_panes = "A2"
    audit.auto_filter.ref = audit.dimensions
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in audit[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
    for col in range(1, audit.max_column + 1):
        width = min(60, max(12, max(len(str(audit.cell(row, col).value or "")) for row in range(1, audit.max_row + 1)) + 2))
        audit.column_dimensions[get_column_letter(col)].width = width

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    wb.save(OUT_FILE)
    print(f"saved={OUT_FILE}")
    print(f"corrected={len(changed)}")
    print(f"review_items={len(notes)}")


if __name__ == "__main__":
    main()
