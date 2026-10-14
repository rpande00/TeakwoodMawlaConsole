from __future__ import annotations

from copy import copy
from datetime import datetime
from pathlib import Path
import json
import re

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(r"D:\Teakwood\Mawla\mawla-booking")
STATE_FILE = ROOT / "data" / "state.json"
OUTPUT_DIR = ROOT / "outputs" / "localhost_matched_excels"

WORKBOOKS = [
    Path(r"D:\Teakwood\Mawla\Mawla Ghaati Group Booking.xlsx"),
    Path(r"D:\Teakwood\Mawla\Mawla Ghaati Group Booking_Chennai.xlsx"),
    Path(r"C:\Users\riyav\Downloads\Mawla Ghati Ultra 2026 (Responses).xlsx"),
]

ACCOMMODATION_EXCEL = {
    "Cottage Dormitory": "Dorm Cottage \u2013 Rs 1500/- per person",
    "Dormitory Hall": "Dormitory Hall \u2013 Rs 1400/- per person",
    "Bamboo House": "Bamboo House \u2013 Rs 1400/- per person",
    "Bring your own tent": "Bring your own Tent - Rs 700/- per spot",
    "3 person tent": "3 person Tent - Complete Tent - Rs 3000/-",
    "4 person tent": "4 person Tent - Complete Tent - Rs 4000/-",
    "8 person tent": "8 person Tent - Rs 1200/-",
}


def norm_name(value):
    return re.sub(r"\s+", " ", re.sub(r"[^a-z0-9]+", " ", str(value or "").casefold()).strip())


def norm_phone(value):
    text = str(value or "").strip()
    if text.endswith(".0"):
        text = text[:-2]
    digits = re.sub(r"\D", "", text)
    if len(digits) > 10 and digits.startswith("91"):
        digits = digits[-10:]
    return digits


def parse_date(value):
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    text = str(value).strip()
    for fmt in ("%m/%d/%Y", "%m/%d/%y", "%Y-%m-%d %H:%M:%S", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    return None


def date_value_for_sheet(source_date, header):
    date = parse_date(source_date)
    if not date:
        return source_date
    if "MM/DD/YYYY" in str(header):
        return f"{date.month}/{date.day}/{date.year}"
    return datetime(date.year, date.month, date.day)


def norm_date(value):
    date = parse_date(value)
    return date.isoformat() if date else str(value or "").strip()


def norm_accommodation(value):
    text = str(value or "").casefold()
    text = re.sub(r"[^a-z0-9]+", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    if "dorm cottage" in text or "cottage dormitory" in text:
        return "Cottage Dormitory"
    if "dormitory hall" in text or "dorm hall" in text:
        return "Dormitory Hall"
    if "bamboo house" in text:
        return "Bamboo House"
    if "bring your own tent" in text:
        return "Bring your own tent"
    if "8 person" in text:
        return "8 person tent"
    if "4 person" in text:
        return "4 person tent"
    if "3 person" in text:
        return "3 person tent"
    return str(value or "").strip()


def norm_nights(value):
    words = {"one": 1, "two": 2, "three": 3, "four": 4}
    if isinstance(value, (int, float)):
        return int(value)
    text = str(value or "").strip().casefold()
    if text in words:
        return words[text]
    try:
        return int(float(text))
    except ValueError:
        return None


def header_map(ws):
    headers = [cell.value for cell in ws[1]]
    normalized = {re.sub(r"[^a-z0-9]+", "", str(h or "").casefold()): idx + 1 for idx, h in enumerate(headers)}
    return headers, normalized


def find_col(normalized, *keys):
    for key in keys:
        compact = re.sub(r"[^a-z0-9]+", "", key.casefold())
        if compact in normalized:
            return normalized[compact]
    for compact, idx in normalized.items():
        if any(compact.startswith(re.sub(r"[^a-z0-9]+", "", key.casefold())) for key in keys):
            return idx
    return None


def build_localhost_index():
    state = json.loads(STATE_FILE.read_text(encoding="utf-8"))
    by_name_phone = {}
    by_name_phone_date = {}
    duplicate_name_phone = set()
    duplicate_name_phone_date = set()
    for booking in state["bookings"]:
        name_phone = (norm_name(booking.get("fullName")), norm_phone(booking.get("phoneNumber")))
        name_phone_date = (*name_phone, norm_date(booking.get("checkInDate")))
        if name_phone in by_name_phone:
            duplicate_name_phone.add(name_phone)
        by_name_phone[name_phone] = booking
        if name_phone_date in by_name_phone_date:
            duplicate_name_phone_date.add(name_phone_date)
        by_name_phone_date[name_phone_date] = booking
    return by_name_phone, by_name_phone_date, duplicate_name_phone, duplicate_name_phone_date, state["bookings"]


def copy_sheet_column_widths(src_ws, dst_ws):
    for key, dim in src_ws.column_dimensions.items():
        dst_ws.column_dimensions[key].width = dim.width


def sync_workbook(path, by_name_phone, by_name_phone_date, duplicate_name_phone, duplicate_name_phone_date):
    wb = openpyxl.load_workbook(path)
    changes = []
    review = []
    matched = 0

    for ws in wb.worksheets:
        if ws.max_row < 2:
            continue
        headers, normalized = header_map(ws)
        name_col = find_col(normalized, "Full Name", "Name", "Guest Name")
        phone_col = find_col(normalized, "Phone Number", "Phone")
        acc_col = find_col(normalized, "Accommodation (per night charges)", "Accommodation Type", "Room Type")
        date_col = find_col(normalized, "Date of Check In (MM/DD/YYYY)", "Date of Check In", "Check-In Date", "CheckInDate")
        nights_col = find_col(normalized, "No of nights of stay : 1", "No of nights of stay", "Nights")
        if not (name_col and phone_col and acc_col and date_col and nights_col):
            continue

        rows_to_delete = []
        for row in range(2, ws.max_row + 1):
            if not any(ws.cell(row, col).value is not None for col in (name_col, phone_col, acc_col, date_col, nights_col)):
                continue
            name_phone = (norm_name(ws.cell(row, name_col).value), norm_phone(ws.cell(row, phone_col).value))
            name_phone_date = (*name_phone, norm_date(ws.cell(row, date_col).value))
            booking = None
            if name_phone_date not in duplicate_name_phone_date:
                booking = by_name_phone_date.get(name_phone_date)
            if booking is None and name_phone not in duplicate_name_phone:
                booking = by_name_phone.get(name_phone)
            if not booking:
                rows_to_delete.append(row)
                review.append([path.name, ws.title, row, ws.cell(row, name_col).value, ws.cell(row, phone_col).value, "No localhost match", "Removed"])
                continue
            matched += 1

            target_date = date_value_for_sheet(booking["checkInDate"], headers[date_col - 1])
            if norm_date(ws.cell(row, date_col).value) != norm_date(target_date):
                old = ws.cell(row, date_col).value
                ws.cell(row, date_col).value = target_date
                if isinstance(target_date, datetime):
                    ws.cell(row, date_col).number_format = "yyyy-mm-dd"
                changes.append([path.name, ws.title, row, ws.cell(row, name_col).value, "Date of Check In", old, target_date])

            target_nights = int(booking["nights"])
            if norm_nights(ws.cell(row, nights_col).value) != target_nights or ws.cell(row, nights_col).value != target_nights:
                old = ws.cell(row, nights_col).value
                ws.cell(row, nights_col).value = target_nights
                changes.append([path.name, ws.title, row, ws.cell(row, name_col).value, "No of nights of stay", old, target_nights])

            target_acc_key = booking["preferredAccommodation"]
            target_acc = ACCOMMODATION_EXCEL.get(target_acc_key, target_acc_key)
            if norm_accommodation(ws.cell(row, acc_col).value) != target_acc_key:
                old = ws.cell(row, acc_col).value
                ws.cell(row, acc_col).value = target_acc
                changes.append([path.name, ws.title, row, ws.cell(row, name_col).value, "Accommodation", old, target_acc])

        for row in sorted(rows_to_delete, reverse=True):
            ws.delete_rows(row, 1)

    if "Localhost Audit" in wb.sheetnames:
        del wb["Localhost Audit"]
    audit = wb.create_sheet("Localhost Audit")
    audit.append(["Section", "Workbook", "Sheet", "Row", "Name", "Field / Issue", "Old / Detail", "New / Status"])
    for item in changes:
        audit.append(["Corrected", *item])
    for workbook, sheet, row, name, _phone, issue, detail in review:
        audit.append(["Review", workbook, sheet, row, name, issue, detail, ""])
    audit.append(["Summary", path.name, "", "", "", "Matched rows", matched, ""])
    audit.append(["Summary", path.name, "", "", "", "Corrections", len(changes), ""])
    audit.append(["Summary", path.name, "", "", "", "Review items", len(review), ""])
    audit.freeze_panes = "A2"
    audit.auto_filter.ref = audit.dimensions
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in audit[1]:
        cell.fill = fill
        cell.font = Font(color="FFFFFF", bold=True)
    for col in range(1, audit.max_column + 1):
        width = min(70, max(12, max(len(str(audit.cell(row, col).value or "")) for row in range(1, audit.max_row + 1)) + 2))
        audit.column_dimensions[get_column_letter(col)].width = width

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    out_path = OUTPUT_DIR / f"{path.stem}_localhost_matched.xlsx"
    wb.save(out_path)
    return out_path, len(changes), len(review), matched


def main():
    by_name_phone, by_name_phone_date, duplicate_name_phone, duplicate_name_phone_date, bookings = build_localhost_index()
    print(f"localhost_bookings={len(bookings)}")
    for path in WORKBOOKS:
        out_path, changes, review, matched = sync_workbook(
            path, by_name_phone, by_name_phone_date, duplicate_name_phone, duplicate_name_phone_date
        )
        print(f"{path.name}|matched={matched}|corrections={changes}|review={review}|output={out_path}")


if __name__ == "__main__":
    main()
